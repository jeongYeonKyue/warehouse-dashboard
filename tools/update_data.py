from pathlib import Path
import json,re
from collections import OrderedDict
from openpyxl import load_workbook

base=Path(__file__).resolve().parents[1]
source=base/'master'/'latest.xlsx'
wb=load_workbook(source,read_only=True,data_only=True)
ws=wb.active
rows=ws.iter_rows(values_only=True)
headers=[str(v or '').replace('\n',' ').strip().lower() for v in next(rows)]
required={'코드','창고','위치','재고수량','상품명','바코드','pcs','ctn','pt용적'}
missing=required-set(headers)
if missing: raise ValueError(f'Excel is missing required columns: {sorted(missing)}')
def val(row,key,default=None):
    try:return row[headers.index(key)] if row[headers.index(key)] is not None else default
    except ValueError:return default
def number(v):
    try:return float(v or 0)
    except (ValueError,TypeError):return 0.0

def parse(code):
    s=str(code or '').strip()
    if s.endswith('.0'):s=s[:-2]
    m=re.fullmatch(r'F([12])(\d{5})',s,re.I)
    if m:
        n=int(m.group(2));start=129 if m.group(1)=='1' and n>=129 else 1
        levels=2 if m.group(1)=='1' and n<=128 else 3
        stack=(n-start)//levels+1
        return 'F'+m.group(1),'01',str(stack).zfill(2),(n-start)%levels+1
    m=re.fullmatch(r'([A-Za-z]\d|\d{2})(\d{2})(\d{2})(\d)',s)
    if m:return m.group(1).upper(),m.group(2),m.group(3),int(m.group(4))
    return None

data=OrderedDict();skipped=0
for row in rows:
    code=str(val(row,'코드','')).strip()
    if code.endswith('.0'):code=code[:-2]
    parts=parse(code)
    if not parts:
        skipped+=1;continue
    z,a,c,t=parts
    warehouse=str(val(row,'창고',''))
    floor=str(val(row,'위치',''))
    key=(warehouse,floor,code)
    if key not in data:
        data[key]={'code':code,'w':warehouse,'floor':floor,'z':z,'a':a,'c':c,'t':t,'use':str(val(row,'용도','')),'type':str(val(row,'구분','')),'qty':0,'pt':0,'items':[]}
    rec=data[key]
    qty=number(val(row,'재고수량'));pt=number(val(row,'pt용적'))
    rec['qty']+=qty;rec['pt']+=pt
    rec['items'].append([str(val(row,'상품명','')),qty,str(val(row,'바코드','')),number(val(row,'pcs')),number(val(row,'ctn')),str(val(row,'상품','')),str(val(row,'연번','')),str(val(row,'색상','')),str(val(row,'묶음','')),str(val(row,'단위','')),pt])
if not data: raise ValueError('No valid location codes found; keeping previous data unchanged')
(base/'latest.json').write_text(json.dumps(list(data.values()),ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print('Updated',len(data),'locations; skipped',skipped,'rows with unsupported codes')
